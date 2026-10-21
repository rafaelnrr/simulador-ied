"""Subida de pontos: lista de sinais que o simulador leva de um estado inicial a um final, para conferir no SAGE.

A lista fica na configuração do IED (cfg["subida"]); a execução roda no servidor, ponto a ponto, e continua
mesmo com o navegador fechado. A planilha modelo (xlsx) traz a lista dos pontos do IED e os estados aceitos,
com listas suspensas, para a importação não ter erro de digitação.
"""
from __future__ import annotations

import asyncio
import csv
import io
import re
import time
import unicodedata


class _IedMod:
    """ied.py importa este módulo; PointError e DPS_TXT são buscados lá só na hora do uso."""

    def __getattr__(self, nome):
        from . import ied
        return getattr(ied, nome)


_ied = _IedMod()

DPS_NOMES = {0: "intermediario", 1: "aberto", 2: "fechado", 3: "invalido"}
BOOL_TXT = {"1": "1", "true": "1", "verdadeiro": "1", "on": "1", "ligado": "1", "sim": "1", "fechado": "1",
            "0": "0", "false": "0", "falso": "0", "off": "0", "desligado": "0", "nao": "0", "aberto": "0"}
FOLHAS = ("$stVal", "$mag$f", "$mag$i", "$cVal$mag$f", "$general", "$actVal")
COLUNAS = {  # cabeçalho normalizado -> campo
    "ponto": "ponto", "sinal": "ponto", "endereco": "ponto", "referencia": "ponto", "caminho": "ponto",
    "estado inicial": "inicial", "inicial": "inicial", "valor inicial": "inicial",
    "estado final": "final", "final": "final", "valor final": "final",
    "observacao": "obs", "obs": "obs", "comentario": "obs",
    "incluir": "incluir", "incluir (s/n)": "incluir", "testar": "incluir",
}
NAO = {"n", "nao", "0", "false", "falso", "x"}
MAX_LINHAS = 5000


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def _txt(v) -> str:
    """Valor de célula como texto (o Excel entrega 1.0, True, ...)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def opcoes(p) -> list[str] | None:
    """Estados que o ponto aceita; None = numérico livre."""
    if p.cdc in ("DPS", "DPC"):
        return ["aberto", "fechado", "intermediario", "invalido"]
    if p.kind == "dig":
        return ["1", "0"]
    if p.v is not None and p.v.enum:
        return [str(t) for t in p.v.enum.values()]
    return None


def aceitos_txt(p) -> str:
    op = opcoes(p)
    if p.cdc in ("DPS", "DPC"):
        return "aberto, fechado, intermediario, invalido"
    if p.kind == "dig":
        return "1 ou 0 (também ligado/desligado, fechado/aberto)"
    if op:
        return ", ".join(op)
    return "número" + (f" de {p.rng[0]:g} a {p.rng[1]:g}" if p.rng else "") + (f" {p.units}" if p.units else "")


def canon(p, valor) -> str:
    """Valor digitado -> forma canônica gravada na lista. ValueError com a lista do que é aceito."""
    s = _txt(valor)
    n = _norm(s)
    if not n:
        raise ValueError("vazio")
    if p.cdc in ("DPS", "DPC"):
        if n.isdigit() and int(n) in DPS_NOMES:
            return DPS_NOMES[int(n)]
        if n in _ied.DPS_TXT:
            return DPS_NOMES[_ied.DPS_TXT[n]]
        raise ValueError(f"use {aceitos_txt(p)}")
    if p.kind == "dig":
        if n in BOOL_TXT:
            return BOOL_TXT[n]
        raise ValueError(f"use {aceitos_txt(p)}")
    op = opcoes(p)
    if op:
        for o in op:
            if _norm(o) == n:
                return o
        if re.fullmatch(r"-?\d+", n):
            for o, t in p.v.enum.items():
                if int(n) == o:
                    return str(t)
        raise ValueError(f"use {aceitos_txt(p)}")
    try:
        f = float(s.replace(",", "."))
    except ValueError:
        raise ValueError(f"use {aceitos_txt(p)}") from None
    return str(int(f)) if f.is_integer() else repr(f)


def endereco_sage(p) -> str:
    return f"{p.ln}${p.fc}${p.do}"


class Subida:
    def __init__(self, ied):
        self.ied = ied
        self.task: asyncio.Task | None = None
        self.pausado = False
        self.parar = False
        self.atual: int | None = None
        self.fila: list[int] = []

    # ------------------------------------------------------------------ dados
    @property
    def cfg(self) -> dict:
        c = self.ied.cfg.setdefault("subida", {})
        c.setdefault("linhas", [])
        c.setdefault("espera_inicial_s", 2.0)
        c.setdefault("espera_entre_s", 3.0)
        c.setdefault("retornar", False)
        c.setdefault("seq", 0)
        return c

    def _salvar(self):
        self.ied.mgr.save(self.ied)

    def linha(self, lid: int) -> dict:
        for ln in self.cfg["linhas"]:
            if ln["id"] == lid:
                return ln
        raise _ied.PointError(f"linha {lid} não existe", status=404)

    def resolver(self, texto: str):
        """Aceita LD/LN$FC$DO, LN$FC$DO, o ID do SAGE (i151-LN$FC$DO), pontos no lugar de $ e o atributo no fim."""
        t = _txt(texto)
        if not t:
            raise ValueError("ponto vazio")
        ied = self.ied
        cands = [t]
        if "." in t:
            cands.append(t.replace(".", "$"))
        for c in list(cands):
            for f in FOLHAS:
                if c.endswith(f):
                    cands.append(c[: -len(f)])
        for c in cands:
            if c in ied.by_path:
                return ied.by_path[c]
        for c in cands:
            ld, _, resto = c.rpartition("/")
            resto = re.sub(r"^[A-Za-z]\w*-", "", resto)  # ID do SAGE: i151-XCBR1$ST$Pos
            achados = [p for p in ied.points if p.ln_path == resto and (not ld or p.domain == ld)]
            if len(achados) == 1:
                return achados[0]
            if len(achados) > 1:
                raise ValueError(f"'{t}' existe em mais de um LD ({', '.join(p.domain for p in achados)}): "
                                 f"escreva LD/{resto}")
        raise ValueError(f"ponto '{t}' não existe neste IED")

    def validar(self, ent: dict) -> tuple[dict, list[str], list[str]]:
        erros, avisos = [], []
        out = {"ponto": "", "inicial": _txt(ent.get("inicial")), "final": _txt(ent.get("final")),
               "obs": _txt(ent.get("obs"))}
        try:
            p = self.resolver(ent.get("ponto"))
        except ValueError as e:
            erros.append(str(e))
            return out, erros, avisos
        out["ponto"] = p.path
        if p.kind == "ro" or p.v is None:
            erros.append(f"{p.path}: CDC {p.cdc} não pode ser alterado pelo simulador")
            return out, erros, avisos
        if p.mode in ("espelhado", "invertido"):
            avisos.append(f"o ponto está no modo {p.mode}: a execução vai recusar; passe para manual ou forçado")
        for campo, nome in (("inicial", "estado inicial"), ("final", "estado final")):
            try:
                out[campo] = canon(p, ent.get(campo))
                self.ied._convert(p, out[campo])
            except (ValueError, _ied.PointError) as e:
                erros.append(f"{nome} '{_txt(ent.get(campo))}' inválido para {p.cdc}: {e}")
        if not erros:
            if out["inicial"] == out["final"]:
                avisos.append("estado inicial igual ao final: o SAGE não vai ver mudança")
            if p.rng and opcoes(p) is None:
                for campo in ("inicial", "final"):
                    v = float(out[campo])
                    if not p.rng[0] <= v <= p.rng[1]:
                        avisos.append(f"{campo} {v:g} fora da faixa do CID [{p.rng[0]:g}, {p.rng[1]:g}]")
        return out, erros, avisos

    def linha_json(self, ln: dict) -> dict:
        p = self.ied.by_path.get(ln["ponto"])
        d = dict(ln)
        if p is None:
            d.update(existe=False, erro_ponto="o ponto não existe mais neste IED (CID reimportado?)")
            return d
        d.update(existe=True, desc=p.desc, cdc=p.cdc, tipo=p.kind, sage=endereco_sage(p), ld=p.domain,
                 valor_atual=self._atual(p), modo=p.mode, opcoes=opcoes(p), aceitos=aceitos_txt(p))
        return d

    def _atual(self, p) -> str:
        v = p.value()
        if p.cdc in ("DPS", "DPC"):
            return DPS_NOMES.get(v, str(v))
        if p.kind == "dig":
            return "1" if v else "0"
        if p.v is not None and p.v.enum and v in p.v.enum:
            return str(p.v.enum[v])
        if isinstance(v, dict):
            return f"{v.get('mag'):g}"
        return f"{v:g}" if isinstance(v, float) else str(v)

    def estado(self) -> dict:
        c = self.cfg
        assocs = self.ied.server.associations() if self.ied.server else []
        return {"linhas": [self.linha_json(ln) for ln in c["linhas"]],
                "config": {k: c[k] for k in ("espera_inicial_s", "espera_entre_s", "retornar")},
                "execucao": self.execucao(),
                "ied": {"estado": self.ied.state, "clientes": [a.peer_str for a in assocs]}}

    def execucao(self) -> dict:
        return {"rodando": self.task is not None and not self.task.done(), "pausado": self.pausado,
                "atual": self.atual, "fila": list(self.fila)}

    def pontos(self) -> list[dict]:
        out = []
        for p in self.ied.points:
            if p.kind == "ro" or p.v is None:
                continue
            out.append({"path": p.path, "sage": endereco_sage(p), "cdc": p.cdc, "tipo": p.kind, "desc": p.desc,
                        "opcoes": opcoes(p), "aceitos": aceitos_txt(p), "valor_atual": self._atual(p),
                        "datasets": sorted({d for d, _ in p.datasets})})
        out.sort(key=lambda x: (x["tipo"] != "dig", x["path"]))
        return out

    def sugestao(self, p) -> tuple[str, str]:
        """Estados sugeridos na planilha preenchida: DPS fechado -> aberto, digital 0 -> 1, enumerado e medida
        partindo do valor atual."""
        if p.cdc in ("DPS", "DPC"):
            return "fechado", "aberto"
        if p.kind == "dig":
            return "0", "1"
        op = opcoes(p)
        atual = self._atual(p)
        if op:
            ini = atual if atual in op else op[0]
            return ini, next((o for o in op if o != ini), ini)
        try:
            cur = float(atual)
        except ValueError:
            cur = 0.0
        if p.rng:
            lo, hi = p.rng
            ini = cur if lo <= cur <= hi else lo
            fim = lo + (hi - lo) / 2
            if fim == ini:
                fim = lo + (hi - lo) * 0.75
        else:
            ini, fim = cur, cur + 1
        return canon(p, round(ini, 3)), canon(p, round(fim, 3))

    def selecionar(self, tipos: set[str], so_dataset: bool) -> list:
        pts = [p for p in self.ied.points if p.kind in tipos and p.v is not None and (p.datasets or not so_dataset)]
        return sorted(pts, key=lambda p: (p.kind != "dig", p.domain, p.ln, p.do))

    # ------------------------------------------------------------------ edição
    def adicionar(self, linhas: list[dict], substituir: bool = False) -> dict:
        if not isinstance(linhas, list) or not linhas:
            raise _ied.PointError("nenhuma linha")
        if self.execucao()["rodando"]:
            raise _ied.PointError("pare a execução antes de mudar a lista", status=409)
        c = self.cfg
        novas, problemas = [], []
        for i, ent in enumerate(linhas, 1):
            out, erros, _ = self.validar(ent)
            if erros:
                problemas.append(f"linha {ent.get('linha', i)}: " + "; ".join(erros))
            novas.append(out)
        if problemas:
            raise _ied.PointError("corrija antes de gravar:\n" + "\n".join(problemas[:50]) +
                             (f"\n… e mais {len(problemas) - 50}" if len(problemas) > 50 else ""), status=422)
        base = [] if substituir else c["linhas"]
        if len(base) + len(novas) > MAX_LINHAS:
            raise _ied.PointError(f"a lista passaria de {MAX_LINHAS} linhas")
        for out in novas:
            c["seq"] += 1
            base.append(dict(out, id=c["seq"], resultado="", nota="", etapa="", executado_em=None,
                             conferido_em=None, erro=""))
        c["linhas"] = base
        self._salvar()
        self._avisar()
        return {"adicionadas": len(novas)}

    def editar(self, lid: int, body: dict) -> dict:
        ln = self.linha(lid)
        if any(k in body for k in ("ponto", "inicial", "final")):
            if self.atual == lid:
                raise _ied.PointError("a linha está em execução", status=409)
            out, erros, _ = self.validar({"ponto": body.get("ponto", ln["ponto"]),
                                          "inicial": body.get("inicial", ln["inicial"]),
                                          "final": body.get("final", ln["final"]), "obs": body.get("obs", ln["obs"])})
            if erros:
                raise _ied.PointError("; ".join(erros), status=422)
            ln.update(out)
        if "obs" in body:
            ln["obs"] = _txt(body["obs"])
        if "resultado" in body:
            if body["resultado"] not in ("", "ok", "falha"):
                raise _ied.PointError("resultado deve ser ok, falha ou vazio")
            ln["resultado"] = body["resultado"]
            ln["conferido_em"] = time.time() if body["resultado"] else None
        if "nota" in body:
            ln["nota"] = _txt(body["nota"])
        self._salvar()
        self._avisar(ln)
        return self.linha_json(ln)

    def remover(self, lid: int):
        if self.atual == lid:
            raise _ied.PointError("a linha está em execução", status=409)
        self.linha(lid)
        self.cfg["linhas"] = [x for x in self.cfg["linhas"] if x["id"] != lid]
        self.fila = [x for x in self.fila if x != lid]
        self._salvar()
        self._avisar()

    def mover(self, lid: int, para: int):
        ls = self.cfg["linhas"]
        ln = self.linha(lid)
        ls.remove(ln)
        ls.insert(max(0, min(para, len(ls))), ln)
        self._salvar()
        self._avisar()

    def apagar_tudo(self):
        if self.execucao()["rodando"]:
            raise _ied.PointError("pare a execução antes de apagar a lista", status=409)
        self.cfg["linhas"] = []
        self._salvar()
        self._avisar()

    def limpar_resultados(self):
        for ln in self.cfg["linhas"]:
            ln.update(resultado="", nota="", etapa="", executado_em=None, conferido_em=None, erro="")
        self._salvar()
        self._avisar()

    def configurar(self, body: dict) -> dict:
        c = self.cfg
        for k in ("espera_inicial_s", "espera_entre_s"):
            if k in body:
                v = float(body[k])
                if not 0 <= v <= 3600:
                    raise _ied.PointError(f"{k} deve ficar entre 0 e 3600 s")
                c[k] = v
        if "retornar" in body:
            c["retornar"] = bool(body["retornar"])
        self._salvar()
        self._avisar()
        return {k: c[k] for k in ("espera_inicial_s", "espera_entre_s", "retornar")}

    # ------------------------------------------------------------------ execução
    def _escrever(self, ln: dict, campo: str):
        p = self.ied.by_path.get(ln["ponto"])
        if p is None:
            raise _ied.PointError("o ponto não existe mais neste IED")
        self.ied.set_points([{"id": p.id, "valor": ln[campo], "confirmar_fora_faixa": True}], origin="subida")

    def aplicar(self, lid: int, etapa: str):
        if self.execucao()["rodando"]:
            raise _ied.PointError("há uma execução em andamento; pause ou pare antes", status=409)
        ln = self.linha(lid)
        if etapa == "comutar":
            self.fila = [lid]
            self.task = asyncio.get_running_loop().create_task(self._rodar())
            return
        if etapa not in ("inicial", "final"):
            raise _ied.PointError("etapa deve ser inicial, final ou comutar")
        try:
            self._escrever(ln, etapa)
            ln.update(etapa=etapa, erro="")
            if etapa == "final":
                ln["executado_em"] = time.time()
        except _ied.PointError as e:
            ln.update(erro=str(e))
            raise
        finally:
            self._salvar()
            self._avisar(ln)

    def executar(self, ids: list[int] | None, so_pendentes: bool):
        if self.execucao()["rodando"]:
            raise _ied.PointError("já há uma execução em andamento", status=409)
        ls = self.cfg["linhas"]
        sel = set(ids or [])
        fila = [x["id"] for x in ls if (not sel or x["id"] in sel)
                and not (so_pendentes and (x["executado_em"] or x["resultado"]))]
        if not fila:
            raise _ied.PointError("nenhuma linha para executar" + (" (todas já executadas ou conferidas)" if so_pendentes else ""))
        self.fila = fila
        self.pausado = self.parar = False
        self.task = asyncio.get_running_loop().create_task(self._rodar())

    def pausar(self, sim: bool):
        self.pausado = sim
        self._avisar()

    def interromper(self):
        self.parar = True
        self.pausado = False
        if self.task and not self.task.done():
            self.task.cancel()

    async def _esperar(self, s: float):
        fim = time.monotonic() + s
        while time.monotonic() < fim or self.pausado:
            await asyncio.sleep(0.1)

    async def _rodar(self):
        c = self.cfg
        try:
            while self.fila:
                while self.pausado:
                    await asyncio.sleep(0.1)
                lid = self.fila[0]
                try:
                    ln = self.linha(lid)
                except _ied.PointError:
                    self.fila.pop(0)
                    continue
                self.atual = lid
                try:
                    ln.update(etapa="inicial", erro="")
                    self._escrever(ln, "inicial")
                    self._avisar(ln)
                    await self._esperar(c["espera_inicial_s"])
                    ln["etapa"] = "final"
                    self._escrever(ln, "final")
                    ln["executado_em"] = time.time()
                    self._avisar(ln)
                    if c["retornar"]:
                        await self._esperar(c["espera_entre_s"])
                        ln["etapa"] = "retornou"
                        self._escrever(ln, "inicial")
                        self._avisar(ln)
                except _ied.PointError as e:
                    ln.update(erro=str(e))
                    self._avisar(ln)
                self.fila.pop(0)
                self.atual = None
                self._salvar()
                if self.fila:
                    await self._esperar(c["espera_entre_s"])
        except asyncio.CancelledError:
            pass
        finally:
            self.atual = None
            self.fila = []
            self.pausado = False
            self._salvar()
            self._avisar()

    def _avisar(self, ln: dict | None = None):
        self.ied.emit("subida", {"linha": self.linha_json(ln) if ln else None, "execucao": self.execucao()})

    # ------------------------------------------------------------------ planilhas
    def ler_arquivo(self, nome: str, data: bytes) -> dict:
        """Lê xlsx ou csv e devolve a prévia: cada linha com o ponto resolvido, erros e avisos."""
        if nome.lower().endswith((".xlsx", ".xlsm")):
            linhas = self._ler_xlsx(data)
        elif nome.lower().endswith((".csv", ".txt")):
            linhas = self._ler_csv(data)
        else:
            raise _ied.PointError("use a planilha modelo (.xlsx) ou um .csv")
        if not linhas:
            raise _ied.PointError("nenhuma linha preenchida na planilha" +
                                  (f" ({self._ignoradas} marcadas com Incluir = N)" if self._ignoradas else ""))
        if len(linhas) > MAX_LINHAS:
            raise _ied.PointError(f"planilha com mais de {MAX_LINHAS} linhas")
        out = []
        for n, ent in linhas:
            norm, erros, avisos = self.validar(ent)
            out.append({"linha": n, "entrada": {k: _txt(ent.get(k)) for k in ("ponto", "inicial", "final", "obs")},
                        **norm, "erros": erros, "avisos": avisos})
        return {"linhas": out, "validas": sum(1 for x in out if not x["erros"]),
                "com_erro": sum(1 for x in out if x["erros"]), "ignoradas": self._ignoradas}

    @staticmethod
    def _mapear(cab: list) -> dict | None:
        m = {}
        for i, c in enumerate(cab):
            k = COLUNAS.get(_norm(c).rstrip(" *"))
            if k and k not in m:
                m[k] = i
        return m if {"ponto", "inicial", "final"} <= set(m) else None

    def _tabela(self, rows) -> list:
        """[(nº da linha, entrada)]; linhas com Incluir = N ficam de fora (contadas em self._ignoradas)."""
        mapa, out = None, []
        self._ignoradas = 0
        for n, row in enumerate(rows, 1):
            row = list(row)
            if mapa is None:
                mapa = self._mapear(row)
                if mapa is None and n >= 15:
                    break
                continue
            ent = {k: row[i] if i < len(row) else None for k, i in mapa.items()}
            if not any(_txt(ent.get(k)) for k in ("ponto", "inicial", "final")):
                continue
            if _norm(ent.pop("incluir", "")) in NAO:
                self._ignoradas += 1
                continue
            out.append((n, ent))
        if mapa is None:
            raise _ied.PointError("cabeçalho não encontrado: a planilha precisa das colunas Ponto, Estado inicial e "
                             "Estado final (baixe a planilha modelo)")
        return out

    def _ler_xlsx(self, data: bytes) -> list:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise _ied.PointError(f"não consegui abrir a planilha: {e}") from e
        ws = wb["Subida"] if "Subida" in wb.sheetnames else wb.worksheets[0]
        return self._tabela(ws.iter_rows(values_only=True))

    def _ler_csv(self, data: bytes) -> list:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        amostra = text[:2000]
        delim = ";" if amostra.count(";") >= amostra.count(",") else ","
        return self._tabela(csv.reader(io.StringIO(text), delimiter=delim))

    def planilha(self, modo: str = "modelo", tipos: set[str] | None = None, so_dataset: bool = True) -> bytes:
        """modo: modelo (300 linhas em branco), preenchida (pontos do IED com estados sugeridos, a revisar)
        ou resultado (a lista atual com a execução e a conferência)."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation

        ied = self.ied
        pts = self.pontos()
        wb = Workbook()
        ws = wb.active
        ws.title = "Subida"
        cab_fill = PatternFill("solid", fgColor="1F4E78")
        auto_fill = PatternFill("solid", fgColor="EDEDED")
        branco = Font(bold=True, color="FFFFFF")
        cab = ["Ponto *", "Estado inicial *", "Estado final *", "Observação", "Incluir (S/N)", "Descrição (automática)",
               "CDC (automático)", "Estados aceitos (automático)"]
        if modo == "resultado":
            cab += ["Executado em", "Resultado no SAGE", "Nota do operador"]
        ws.append(cab)
        for c in ws[1]:
            c.fill, c.font = cab_fill, branco
            c.alignment = Alignment(wrap_text=True, vertical="center")
        n_pts = len(pts) + 1
        if modo == "resultado":
            linhas = self.cfg["linhas"]
        elif modo == "preenchida":
            linhas = []
            for p in self.selecionar(tipos or {"dig"}, so_dataset):
                ini, fim = self.sugestao(p)
                linhas.append({"ponto": p.path, "inicial": ini, "final": fim, "obs": ""})
        else:
            linhas = []
        total = len(linhas) if modo != "modelo" else 300
        for r in range(2, total + 2):
            ln = linhas[r - 2] if r - 2 < len(linhas) else None
            ws.cell(r, 1, ln["ponto"] if ln else None)
            ws.cell(r, 2, ln["inicial"] if ln else None)
            ws.cell(r, 3, ln["final"] if ln else None)
            ws.cell(r, 4, ln["obs"] if ln else None)
            ws.cell(r, 5, "S" if ln else None)
            for col, k in ((6, 5), (7, 3), (8, 6)):
                c = ws.cell(r, col, f'=IF(A{r}="","",IFERROR(VLOOKUP(A{r},Pontos!$A$2:$F${n_pts},{k},FALSE),'
                                    f'"ponto não encontrado: escolha na lista"))')
                c.fill = auto_fill
            if ln and modo == "resultado":
                ws.cell(r, 9, time.strftime("%d/%m/%Y %H:%M:%S", time.localtime(ln["executado_em"]))
                        if ln["executado_em"] else "")
                ws.cell(r, 10, {"ok": "OK", "falha": "FALHA"}.get(ln["resultado"], ""))
                ws.cell(r, 11, ln["nota"])
        ultima = max(total + 1, 2)
        dv_p = DataValidation(type="list", formula1=f"=Pontos!$A$2:$A${n_pts}", allow_blank=True,
                              showErrorMessage=True, errorTitle="Ponto inválido",
                              error="Escolha um ponto da aba Pontos (lista suspensa).")
        dv_e = DataValidation(type="list", formula1='"fechado,aberto,intermediario,invalido,1,0"', allow_blank=True,
                              showErrorMessage=True, errorStyle="warning", errorTitle="Estado fora da lista",
                              error="Para medidas e enumerados digite o valor; confira a coluna Estados aceitos.")
        dv_i = DataValidation(type="list", formula1='"S,N"', allow_blank=True, showErrorMessage=True,
                              errorTitle="Incluir", error="S = testar este sinal, N = deixar de fora.")
        for dv in (dv_p, dv_e, dv_i):
            ws.add_data_validation(dv)
        dv_p.add(f"A2:A{ultima}")
        dv_e.add(f"B2:C{ultima}")
        dv_i.add(f"E2:E{ultima}")
        for col, w in zip("ABCDEFGHIJK", (52, 16, 16, 30, 9, 40, 10, 40, 20, 18, 30)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{chr(64 + len(cab))}{ultima}"

        wp = wb.create_sheet("Pontos")
        wp.append(["Ponto", "Endereço no SAGE (LN$FC$DO)", "CDC", "Tipo", "Descrição", "Estados aceitos", "Valor atual",
                   "DataSets"])
        for c in wp[1]:
            c.fill, c.font = cab_fill, branco
        for p in pts:
            wp.append([p["path"], p["sage"], p["cdc"], {"dig": "digital", "ana": "analógico"}.get(p["tipo"], p["tipo"]),
                       p["desc"], p["aceitos"], p["valor_atual"], ", ".join(p["datasets"])])
        for col, w in zip("ABCDEFGH", (52, 32, 8, 11, 40, 44, 14, 50)):
            wp.column_dimensions[col].width = w
        wp.freeze_panes = "A2"
        wp.auto_filter.ref = f"A1:H{n_pts}"

        wi = wb.create_sheet("Instruções")
        for linha in [
            f"Subida de pontos — IED {ied.name} ({ied.ied_name})",
            "",
            "1. Preencha a aba Subida, uma linha por sinal. Colunas com * são obrigatórias.",
            "2. Ponto: escolha na lista suspensa (vem da aba Pontos). Também são aceitos o endereço sem o LD "
            "(XCBR1$ST$Pos), o ID do SAGE (i151-XCBR1$ST$Pos) e o atributo no fim ($stVal).",
            "3. Estado inicial e final: DPS/DPC aceitam aberto, fechado, intermediario, invalido (ou 0 a 3); "
            "SPS e outros digitais aceitam 1/0, ligado/desligado ou fechado/aberto; medidas aceitam número; "
            "enumerados aceitam o texto ou o número. A coluna 'Estados aceitos' mostra o que vale para cada ponto.",
            "4. Incluir: S (ou vazio) testa o sinal; N deixa a linha de fora na importação sem precisar apagá-la.",
            "5. As colunas cinzas são automáticas: se aparecer 'ponto não encontrado', o ponto está errado.",
            "Na planilha preenchida, os estados são sugestões (DPS fechado → aberto, digitais 0 → 1, medidas a partir "
            "do valor atual): revise antes de importar.",
            "6. No simulador: aba Subida de pontos › Importar planilha. A prévia mostra cada linha com erro "
            "antes de gravar.",
            "7. Na execução, para cada linha o simulador aplica o estado inicial, espera, aplica o final e "
            "passa para a próxima. Marque no simulador o que o SAGE mostrou (OK ou Falha).",
        ]:
            wi.append([linha])
        wi["A1"].font = Font(bold=True, size=13)
        wi.column_dimensions["A"].width = 140
        wb.active = 0
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
