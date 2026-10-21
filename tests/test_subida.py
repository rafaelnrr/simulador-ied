"""Subida de pontos: validação da lista, importação de planilha (xlsx e csv), planilha modelo e execução."""
from __future__ import annotations

import io
import time

from conftest import CTRL, MED
from openpyxl import Workbook, load_workbook

POS = f"{CTRL}/XCBR1$ST$Pos"
IND3 = f"{CTRL}/sAutGGIO2$ST$Ind3"
HZ = f"{MED}/MMXU1$MX$Hz"


def _xlsx(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Subida"
    ws.append(["Título qualquer"])
    ws.append(["Ponto *", "Estado inicial *", "Estado final *", "Observação"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _esperar(sim, iid, cond, timeout=5.0):
    fim = time.time() + timeout
    while time.time() < fim:
        d = sim.h.get(f"/ieds/{iid}/subida").json()
        if cond(d):
            return d
        time.sleep(0.1)
    raise AssertionError(d)


def test_lista_manual_validacao(sim):
    iid = sim.create(start=False)["id"]
    u = f"/ieds/{iid}/subida"
    r = sim.h.post(u + "/linhas", json={"linhas": [
        {"ponto": "i151-XCBR1$ST$Pos", "inicial": "Fechado", "final": "0x"},
        {"ponto": "GGIO9$ST$Ind1", "inicial": "1", "final": "0"}]})
    assert r.status_code == 422
    msg = r.json()["erro"] if "erro" in r.json() else r.text
    assert "linha 1" in msg and "estado final" in msg and "linha 2" in msg and "não existe" in msg
    r = sim.h.post(u + "/linhas", json={"linhas": [
        {"ponto": "i151-XCBR1$ST$Pos", "inicial": "Fechado", "final": "1"},
        {"ponto": "sAutGGIO2.ST.Ind3$stVal", "inicial": "fechado", "final": "desligado", "obs": "mola"},
        {"ponto": HZ, "inicial": "60", "final": "59,5"}]})
    assert r.status_code == 200, r.text
    d = sim.h.get(u).json()
    assert [(x["ponto"], x["inicial"], x["final"]) for x in d["linhas"]] == [
        (POS, "fechado", "aberto"), (IND3, "1", "0"), (HZ, "60", "59.5")]
    assert d["linhas"][0]["opcoes"] == ["aberto", "fechado", "intermediario", "invalido"]
    assert d["linhas"][0]["sage"] == "XCBR1$ST$Pos"
    lid = d["linhas"][2]["id"]
    assert sim.h.patch(f"{u}/linhas/{lid}", json={"final": "abc"}).status_code == 422
    assert sim.h.patch(f"{u}/linhas/{lid}", json={"posicao": 0}).status_code == 200
    assert sim.h.get(u).json()["linhas"][0]["ponto"] == HZ
    sim.h.delete(u)
    assert sim.h.get(u).json()["linhas"] == []


def test_planilha_modelo_e_importacao(sim):
    iid = sim.create(start=False)["id"]
    r = sim.h.get(f"/ieds/{iid}/subida-planilha")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Subida", "Pontos", "Instruções"]
    pts = [row[0] for row in wb["Pontos"].iter_rows(min_row=2, values_only=True)]
    assert POS in pts and HZ in pts
    assert wb["Subida"].data_validations.dataValidation  # listas suspensas

    arq = _xlsx([[POS, "aberto", "fechado", "disjuntor"], [None, None, None], ["XCBR1$ST$Pos", 1.0, 2.0],
                 ["NAO_EXISTE", "1", "0"], [IND3, True, False], [HZ, "60", "60"]])
    r = sim.h.post(f"/ieds/{iid}/subida-importar", files={"arquivo": ("lista.xlsx", arq)})
    assert r.status_code == 200, r.text
    p = r.json()
    assert (p["validas"], p["com_erro"]) == (4, 1)
    por_linha = {x["linha"]: x for x in p["linhas"]}
    assert por_linha[3]["ponto"] == POS and por_linha[3]["obs"] == "disjuntor"
    assert (por_linha[5]["inicial"], por_linha[5]["final"]) == ("aberto", "fechado")
    assert por_linha[6]["erros"] and por_linha[7]["final"] == "0"
    assert any("igual" in a for a in por_linha[8]["avisos"])

    csv_txt = "ponto;estado inicial;estado final\nXCBR1$ST$Pos;fechado;aberto\n"
    r = sim.h.post(f"/ieds/{iid}/subida-importar", files={"arquivo": ("l.csv", csv_txt.encode())})
    assert r.json()["validas"] == 1

    validas = [x for x in p["linhas"] if not x["erros"]]
    assert sim.h.post(f"/ieds/{iid}/subida/linhas", json={"linhas": validas, "modo": "substituir"}).status_code == 200
    r = sim.h.get(f"/ieds/{iid}/subida-planilha", params={"resultado": True})
    ws = load_workbook(io.BytesIO(r.content))["Subida"]
    assert [c.value for c in ws[1]][-3:] == ["Executado em", "Resultado no SAGE", "Nota do operador"]
    assert ws.cell(2, 1).value == POS
    # a planilha de resultado volta a ser importável
    r = sim.h.post(f"/ieds/{iid}/subida-importar", files={"arquivo": ("r.xlsx", r.content)})
    assert r.json()["validas"] == 4


def test_execucao(sim):
    iid = sim.create(start=False)["id"]
    u = f"/ieds/{iid}/subida"
    sim.h.put(u + "/config", json={"espera_inicial_s": 0.2, "espera_entre_s": 0.1, "retornar": False})
    sim.h.post(u + "/linhas", json={"linhas": [
        {"ponto": POS, "inicial": "fechado", "final": "aberto"},
        {"ponto": IND3, "inicial": "0", "final": "1"}]})
    r = sim.h.post(u + "/executar", json={})
    assert r.status_code == 200 and r.json()["rodando"]
    d = _esperar(sim, iid, lambda d: not d["execucao"]["rodando"])
    assert all(x["executado_em"] and x["etapa"] == "final" for x in d["linhas"])
    assert [x["valor_atual"] for x in d["linhas"]] == ["aberto", "1"]
    hist = sim.h.get(f"/ieds/{iid}/history", params={"ponto": "XCBR1"}).json()["mudancas"]
    assert [h["valor"] for h in hist[:2]] == [1, 2] and hist[0]["origem"] == "subida"

    # só pendentes: nada a fazer
    assert sim.h.post(u + "/executar", json={}).status_code == 400
    lid = d["linhas"][0]["id"]
    assert sim.h.patch(f"{u}/linhas/{lid}", json={"resultado": "ok", "nota": "SAGE viu"}).json()["conferido_em"]
    r = sim.h.post(f"{u}/linhas/{lid}/aplicar", json={"etapa": "inicial"})
    assert r.json()["valor_atual"] == "fechado"

    # pausar e parar
    sim.h.put(u + "/config", json={"espera_inicial_s": 5})
    sim.h.post(u + "/executar", json={"so_pendentes": False})
    assert sim.h.post(u + "/pausar").json()["pausado"]
    assert sim.h.post(u + "/linhas", json={"linhas": [{"ponto": POS, "inicial": "1", "final": "2"}]}).status_code == 409
    sim.h.post(u + "/parar")
    _esperar(sim, iid, lambda d: not d["execucao"]["rodando"])
    sim.h.delete(u, params={"resultados": True})
    assert not any(x["resultado"] or x["executado_em"] for x in sim.h.get(u).json()["linhas"])


def test_planilha_preenchida_e_coluna_incluir(sim):
    iid = sim.create(start=False)["id"]
    r = sim.h.get(f"/ieds/{iid}/subida-planilha", params={"modo": "preenchida", "tipos": "dig", "so_dataset": False})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content))["Subida"]
    linhas = {row[0]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    assert linhas[POS][1:5] == ("fechado", "aberto", None, "S")
    assert linhas[IND3][1:3] == ("0", "1")
    assert HZ not in linhas
    so_ds = load_workbook(io.BytesIO(sim.h.get(f"/ieds/{iid}/subida-planilha",
                                               params={"modo": "preenchida", "tipos": "dig,ana"}).content))["Subida"]
    pts = [row[0] for row in so_ds.iter_rows(min_row=2, values_only=True)]
    assert HZ in pts and len(pts) <= ws.max_row - 1 + 50

    # Incluir = N fica de fora; o arquivo baixado volta inteiro na importação
    for row in ws.iter_rows(min_row=2):
        if row[0].value != POS:
            row[4].value = "N"
    buf = io.BytesIO()
    ws.parent.save(buf)
    p = sim.h.post(f"/ieds/{iid}/subida-importar", files={"arquivo": ("p.xlsx", buf.getvalue())}).json()
    assert p["validas"] == 1 and p["ignoradas"] == ws.max_row - 2
    assert p["linhas"][0]["ponto"] == POS
